from app.core.intelligence import extract_clues, infer_category


def extract_text(path):
    """Extract text/metadata.

    Extraction errors are intentionally raised to the scanner. The scanner
    catalogs the file first and records the extraction failure on the document.
    """
    ext = path.suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return (
            "\n".join((page.extract_text() or "") for page in reader.pages),
            len(reader.pages), 0, 0,
        )

    if ext == ".docx":
        from docx import Document
        document = Document(str(path))
        out = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                out.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(out), 0, 0, 0

    if ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(
            str(path), read_only=True, data_only=True
        )
        out = []
        try:
            for sheet in workbook.worksheets:
                out.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    out.append(
                        " | ".join("" if value is None else str(value)
                                   for value in row)
                    )
        finally:
            workbook.close()
        return "\n".join(out), 0, 0, 0

    if ext == ".pptx":
        from pptx import Presentation
        presentation = Presentation(str(path))
        out = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    out.append(shape.text)
        return "\n".join(out), len(presentation.slides), 0, 0

    if ext in {".txt", ".md", ".markdown", ".csv", ".html", ".htm"}:
        return path.read_text(
            encoding="utf-8", errors="ignore"
        ), 0, 0, 0

    if ext in {".jpg", ".jpeg", ".png", ".bmp", ".gif",
               ".tif", ".tiff", ".webp"}:
        from PIL import Image
        with Image.open(path) as image:
            return "", 0, image.width, image.height

    return "", 0, 0, 0


def render_pdf_page(path, page_number, zoom=1.25):
    import fitz
    document = fitz.open(str(path))
    try:
        if page_number < 0 or page_number >= len(document):
            return None
        page = document.load_page(page_number)
        pix = page.get_pixmap(
            matrix=fitz.Matrix(zoom, zoom), alpha=False
        )
        return pix.tobytes("png")
    finally:
        document.close()
